from datetime import datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from bd import obtenerconexion
from controlador_tickets import SQL_TICKETS, _completar
from utilidades import AREAS, CATEGORIAS, ESTADOS, PRIORIDADES

COLUMNAS = ["Código", "Asunto", "Categoría", "Prioridad", "Estado", "Solicitante", "Empresa", "Técnico",
            "Creado", "Tiempo (SLA)"]


def filtros_desde(args):
    """Lee los filtros de la URL. Por defecto, los últimos 30 días."""
    hoy = datetime.now().date()
    return {
        "desde": args.get("desde") or (hoy - timedelta(days=30)).isoformat(),
        "hasta": args.get("hasta") or hoy.isoformat(),
        "area": args.get("area") if args.get("area") in AREAS else "",
        "tecnico": args.get("tecnico") or "",
    }


def _where(usuario, empresa, filtros):
    condiciones, parametros = [], []
    if usuario["rol"] == "Solicitante":
        condiciones.append("t.solicitante_id = %s")
        parametros.append(usuario["id"])
    if empresa:
        condiciones.append("t.empresa = %s")
        parametros.append(empresa)
    for clave, hora, operador in (("desde", time.min, ">="), ("hasta", time.max, "<=")):
        try:
            fecha = datetime.combine(datetime.strptime(filtros[clave], "%Y-%m-%d").date(), hora)
            condiciones.append("t.created_at " + operador + " %s")
            parametros.append(fecha)
        except (TypeError, ValueError):
            pass
    if filtros["area"]:
        condiciones.append("t.solicitante_id IN (SELECT id FROM users WHERE area = %s)")
        parametros.append(filtros["area"])
    if filtros["tecnico"] == "sin":
        condiciones.append("t.tecnico_id IS NULL")
    elif filtros["tecnico"]:
        condiciones.append("t.tecnico_id = %s")
        parametros.append(filtros["tecnico"])
    where = " WHERE " + " AND ".join(condiciones) if condiciones else ""
    return where, parametros


def tickets_reporte(usuario, empresa, filtros):
    where, parametros = _where(usuario, empresa, filtros)
    conexion = obtenerconexion()
    with conexion:
        with conexion.cursor() as cursor:
            cursor.execute(SQL_TICKETS + where + " ORDER BY t.created_at DESC", parametros)
            return [_completar(t) for t in cursor.fetchall()]


def resumen(usuario, empresa, filtros):
    """Datos del dashboard. Los conteos se hacen con GROUP BY en MySQL."""
    where, parametros = _where(usuario, empresa, filtros)
    conexion = obtenerconexion()
    with conexion:
        with conexion.cursor() as cursor:
            def agrupar(columna):
                cursor.execute("SELECT t." + columna + " AS clave, COUNT(*) AS total FROM tickets t" + where +
                               " GROUP BY t." + columna, parametros)
                return {fila["clave"]: fila["total"] for fila in cursor.fetchall()}

            por_estado = agrupar("estado")
            por_prioridad = agrupar("prioridad")
            por_categoria = agrupar("categoria")

            cursor.execute("""
                SELECT t.tecnico_id,
                       COUNT(*) AS asignados,
                       SUM(CASE WHEN t.estado = 'En proceso' THEN 1 ELSE 0 END) AS en_proceso,
                       SUM(CASE WHEN t.estado IN ('Resuelto', 'Cerrado') THEN 1 ELSE 0 END) AS resueltos
                FROM tickets t""" + where + " GROUP BY t.tecnico_id", parametros)
            por_tecnico = {fila["tecnico_id"]: fila for fila in cursor.fetchall()}

            cursor.execute("SELECT id, nombre FROM users WHERE rol = 'Técnico' ORDER BY nombre")
            tecnicos = cursor.fetchall()

    # El semáforo depende de la hora actual, por eso se calcula en Python
    tickets = tickets_reporte(usuario, empresa, filtros)
    sla = {"verde": 0, "amarillo": 0, "rojo": 0}
    verdes_por_tecnico = {}
    for t in tickets:
        sla[t["sla"]["nivel"]] += 1
        if t["sla"]["nivel"] == "verde":
            verdes_por_tecnico[t["tecnico_id"]] = verdes_por_tecnico.get(t["tecnico_id"], 0) + 1

    responsables = []
    for tec in tecnicos:
        fila = por_tecnico.get(tec["id"], {"asignados": 0, "en_proceso": 0, "resueltos": 0})
        asignados = int(fila["asignados"])
        responsables.append({
            "id": tec["id"], "nombre": tec["nombre"], "asignados": asignados,
            "en_proceso": int(fila["en_proceso"] or 0), "resueltos": int(fila["resueltos"] or 0),
            "a_tiempo": round(verdes_por_tecnico.get(tec["id"], 0) / asignados * 100) if asignados else None,
        })

    total = len(tickets)
    return {
        "total": total,
        "por_estado": {e: por_estado.get(e, 0) for e in ESTADOS},
        "por_prioridad": {p: por_prioridad.get(p, 0) for p in PRIORIDADES},
        "por_categoria": {c: por_categoria.get(c, 0) for c in CATEGORIAS},
        "sla": sla,
        "sla_porcentaje": round(sla["verde"] / total * 100) if total else None,
        "responsables": responsables,
        "sin_asignar": int(por_tecnico[None]["asignados"]) if None in por_tecnico else 0,
    }


def descripcion_filtros(filtros, empresa):
    partes = ["Del " + filtros["desde"] + " al " + filtros["hasta"], empresa or "Todas las empresas"]
    if filtros["area"]:
        partes.append("Área: " + filtros["area"])
    return " · ".join(partes)


def _filas(tickets):
    return [[t["codigo"], t["asunto"], t["categoria"], t["prioridad"], t["estado"], t["solicitante_nombre"],
             t["empresa"], t["tecnico_nombre"] or "Sin asignar", t["created_at"].strftime("%d/%m/%Y %H:%M"),
             t["sla"]["texto"]] for t in tickets]


def generar_excel(tickets, subtitulo):
    """Excel con openpyxl."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tickets"
    hoja["A1"] = "Reporte de tickets · Mesa de Ayuda TI"
    hoja["A1"].font = Font(bold=True, size=14, color="3F4C8C")
    hoja["A2"] = subtitulo
    hoja["A2"].font = Font(italic=True, color="6B707A")
    hoja.append([])
    hoja.append(COLUMNAS)
    for celda in hoja[4]:
        celda.font = Font(bold=True, color="3F4C8C")
        celda.fill = PatternFill("solid", fgColor="D9DEF5")
        celda.alignment = Alignment(vertical="center")
    for fila in _filas(tickets):
        hoja.append(fila)
    for i, ancho in enumerate([10, 40, 18, 10, 12, 22, 20, 20, 17, 22], start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A5"
    if tickets:
        hoja.auto_filter.ref = "A4:J" + str(hoja.max_row)
    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def generar_pdf(tickets, subtitulo, datos):
    """PDF con ReportLab."""
    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                                  topMargin=14 * mm, bottomMargin=14 * mm, title="Reporte de tickets")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], textColor=colors.HexColor("#3F4C8C"), alignment=0)
    sub = ParagraphStyle("sub", parent=estilos["Normal"], textColor=colors.HexColor("#6B707A"))
    celda = ParagraphStyle("celda", parent=estilos["Normal"], fontSize=8, leading=10)

    estados = datos["por_estado"]
    porcentaje = str(datos["sla_porcentaje"]) + "%" if datos["sla_porcentaje"] is not None else "—"
    resumen_tabla = Table([["Total", "Abiertos", "En proceso", "Resueltos", "Cerrados", "Dentro del plazo"],
                           [datos["total"], estados["Abierto"], estados["En proceso"], estados["Resuelto"],
                            estados["Cerrado"], porcentaje]])
    resumen_tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9DEF5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#3F4C8C")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E8E3DA")),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elementos = [Paragraph("Reporte de tickets · Mesa de Ayuda TI", titulo), Paragraph(subtitulo, sub),
                 Spacer(1, 6 * mm), resumen_tabla, Spacer(1, 6 * mm)]

    if tickets:
        filas = [COLUMNAS] + [[Paragraph(str(v), celda) for v in fila] for fila in _filas(tickets)]
        tabla = Table(filas, repeatRows=1, colWidths=[18 * mm, 58 * mm, 26 * mm, 16 * mm, 20 * mm, 28 * mm,
                                                      28 * mm, 28 * mm, 25 * mm, 28 * mm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3F4C8C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F4EE")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E8E3DA")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("No hay tickets en el rango elegido.", sub))

    documento.build(elementos)
    salida.seek(0)
    return salida
