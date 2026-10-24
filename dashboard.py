"""
MÓDULO 6 — Dashboard, reportes (Excel / PDF) y resumen por API.

* Consultas GROUP BY por estado, prioridad, tipo (categoría) y responsable.
* Filtros por fecha, área y responsable (y empresa para el superusuario).
* Cumplimiento del plazo calculado con el SLA de Python (Módulo 4).
* Exporta lo filtrado a Excel (openpyxl) y PDF (ReportLab).
"""
import io
from datetime import datetime, timedelta

from flask import Blueprint, g, jsonify, render_template, request, send_file

from auth import empresa_filtro, es_super, login_requerido
from db import consultar
from tickets import SQL_TICKETS, alcance, preparar
from utils import ESTADOS, ahora, entero

dashboard_bp = Blueprint("dashboard", __name__)


def _fecha(texto):
    try:
        return datetime.strptime(texto, "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def leer_filtros():
    hoy = ahora().date()
    desde = _fecha(request.args.get("desde")) or datetime.combine(hoy - timedelta(days=30), datetime.min.time())
    hasta = _fecha(request.args.get("hasta")) or datetime.combine(hoy, datetime.min.time())
    if hasta < desde:
        desde, hasta = hasta, desde
    return {
        "desde": desde,
        "hasta": hasta,
        "area": (request.args.get("area") or "").strip(),
        "responsable": request.args.get("responsable") or "",
        "empresa_id": entero(request.args.get("empresa_id")) if es_super() and request.args.get("empresa_id") else empresa_filtro(),
    }


def _where(filtros):
    condicion, params = alcance(g.usuario, filtros["empresa_id"])
    condiciones = [condicion, "t.created_at >= %s", "t.created_at < %s"]
    params += [filtros["desde"], filtros["hasta"] + timedelta(days=1)]
    if filtros["area"]:
        condiciones.append("a.nombre = %s")
        params.append(filtros["area"])
    if filtros["responsable"] == "sin":
        condiciones.append("t.tecnico_id IS NULL")
    elif entero(filtros["responsable"]):
        condiciones.append("t.tecnico_id = %s")
        params.append(entero(filtros["responsable"]))
    return " AND ".join(condiciones), params


def datos_dashboard(filtros):
    where, params = _where(filtros)
    base = """FROM tickets t
              JOIN empresas e    ON e.id = t.empresa_id
              LEFT JOIN areas a  ON a.id = t.area_id
              JOIN prioridades p ON p.id = t.prioridad_id
              JOIN categorias c  ON c.id = t.categoria_id
              LEFT JOIN users tec ON tec.id = t.tecnico_id
             WHERE """ + where

    por_estado = {e: 0 for e in ESTADOS}
    for f in consultar(f"SELECT t.estado, COUNT(*) AS total {base} GROUP BY t.estado", params):
        por_estado[f["estado"]] = f["total"]

    por_prioridad = consultar(
        f"""SELECT p.nombre, COUNT(t.id) AS total
              FROM prioridades p
              LEFT JOIN (SELECT t.prioridad_id, t.id {base}) t ON t.prioridad_id = p.id
             GROUP BY p.id, p.nombre, p.nivel ORDER BY p.nivel DESC""",
        params,
    )
    por_tipo = consultar(
        f"SELECT c.nombre, COUNT(*) AS total {base} GROUP BY c.id, c.nombre ORDER BY total DESC, c.nombre", params
    )
    responsables = consultar(
        f"""SELECT t.tecnico_id, COALESCE(tec.nombre, 'Sin asignar') AS tecnico,
                   COUNT(*) AS asignados,
                   SUM(t.estado = 'En proceso') AS en_proceso,
                   SUM(t.estado IN ('Resuelto', 'Cerrado')) AS resueltos,
                   SUM(t.estado IN ('Resuelto', 'Cerrado')
                       AND COALESCE(t.fecha_resolucion, t.fecha_cierre) <= t.fecha_limite) AS a_tiempo
              {base}
             GROUP BY t.tecnico_id, tec.nombre
             ORDER BY t.tecnico_id IS NULL, asignados DESC""",
        params,
    )
    for r in responsables:
        r["en_proceso"], r["resueltos"], r["a_tiempo"] = int(r["en_proceso"] or 0), int(r["resueltos"] or 0), int(r["a_tiempo"] or 0)
        r["porcentaje"] = round(100 * r["a_tiempo"] / r["resueltos"]) if r["resueltos"] else None

    # SLA (Python): cumplido cuenta como verde, incumplido como rojo
    momento = ahora()
    tickets = [preparar(t, momento) for t in consultar(SQL_TICKETS + " WHERE " + where + " ORDER BY t.created_at DESC", params)]
    sla = {"verde": 0, "amarillo": 0, "rojo": 0}
    for t in tickets:
        clave = {"cumplido": "verde", "incumplido": "rojo"}.get(t["semaforo"], t["semaforo"])
        sla[clave] += 1
    total = len(tickets)
    return {
        "total": total,
        "por_estado": por_estado,
        "por_prioridad": por_prioridad,
        "por_tipo": por_tipo,
        "responsables": responsables,
        "sla": sla,
        "cumplimiento": round(100 * sla["verde"] / total) if total else 0,
        "tickets": tickets,
    }


def opciones_filtros(filtros):
    empresa_id = filtros["empresa_id"]
    if empresa_id:
        areas = consultar("SELECT nombre FROM areas WHERE empresa_id = %s AND activo = 1 ORDER BY nombre", (empresa_id,))
        tecnicos = consultar("SELECT id, nombre FROM users WHERE empresa_id = %s AND rol = 'Técnico' ORDER BY nombre", (empresa_id,))
    else:
        areas = consultar("SELECT DISTINCT nombre FROM areas WHERE activo = 1 ORDER BY nombre")
        tecnicos = consultar("SELECT id, nombre FROM users WHERE rol = 'Técnico' ORDER BY nombre")
    return [a["nombre"] for a in areas], tecnicos


def texto_filtros(filtros):
    partes = [f"Del {filtros['desde']:%d/%m/%Y} al {filtros['hasta']:%d/%m/%Y}"]
    if filtros["empresa_id"]:
        e = consultar("SELECT nombre FROM empresas WHERE id = %s", (filtros["empresa_id"],), uno=True)
        partes.append(f"Empresa: {e['nombre'] if e else '-'}")
    else:
        partes.append("Empresa: todas")
    if filtros["area"]:
        partes.append(f"Área: {filtros['area']}")
    if filtros["responsable"] == "sin":
        partes.append("Responsable: sin asignar")
    elif entero(filtros["responsable"]):
        r = consultar("SELECT nombre FROM users WHERE id = %s", (entero(filtros["responsable"]),), uno=True)
        partes.append(f"Responsable: {r['nombre'] if r else '-'}")
    return " | ".join(partes)


@dashboard_bp.get("/dashboard")
@login_requerido
def dashboard():
    filtros = leer_filtros()
    datos = datos_dashboard(filtros)
    areas, tecnicos = opciones_filtros(filtros)
    return render_template("dashboard.html", d=datos, filtros=filtros, areas=areas, tecnicos=tecnicos,
                           query=request.query_string.decode())


# ------------------------------------------------------------------ Excel
@dashboard_bp.get("/dashboard/exportar/excel")
@login_requerido
def exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    filtros = leer_filtros()
    d = datos_dashboard(filtros)
    color = (g.usuario.get("color_primario") or "#3D4A8A").lstrip("#")
    cabecera = PatternFill("solid", fgColor=color)
    blanco = Font(bold=True, color="FFFFFF")

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    hoja["A1"] = "Reporte de tickets - Mesa de Ayuda TI"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = texto_filtros(filtros)
    hoja["A3"] = f"Generado: {ahora():%d/%m/%Y %H:%M} por {g.usuario['nombre']}"
    fila = 5

    def bloque(titulo, encabezados, filas):
        nonlocal fila
        hoja.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12)
        fila += 1
        for col, texto in enumerate(encabezados, 1):
            celda = hoja.cell(row=fila, column=col, value=texto)
            celda.fill, celda.font = cabecera, blanco
        for valores in filas:
            fila += 1
            for col, valor in enumerate(valores, 1):
                hoja.cell(row=fila, column=col, value=valor)
        fila += 2

    bloque("Por estado", ["Estado", "Tickets"], list(d["por_estado"].items()))
    bloque("Por prioridad", ["Prioridad", "Tickets"], [(p["nombre"], p["total"]) for p in d["por_prioridad"]])
    bloque("Por tipo", ["Categoría", "Tickets"], [(c["nombre"], c["total"]) for c in d["por_tipo"]])
    bloque("Cumplimiento del plazo", ["Semáforo", "Tickets"],
           [("Verde (a tiempo)", d["sla"]["verde"]), ("Amarillo (por vencer)", d["sla"]["amarillo"]),
            ("Rojo (vencido)", d["sla"]["rojo"]), ("% dentro del plazo", f"{d['cumplimiento']}%")])
    bloque("Por responsable", ["Técnico", "Asignados", "En proceso", "Resueltos", "A tiempo"],
           [(r["tecnico"], r["asignados"], r["en_proceso"], r["resueltos"],
             f"{r['porcentaje']}%" if r["porcentaje"] is not None else "-") for r in d["responsables"]])
    hoja.column_dimensions["A"].width = 32
    for col in "BCDE":
        hoja.column_dimensions[col].width = 14

    detalle = libro.create_sheet("Tickets")
    columnas = ["Código", "Empresa", "Asunto", "Categoría", "Prioridad", "Estado", "Solicitante", "Área",
                "Técnico", "Creado", "Fecha límite", "SLA"]
    for col, texto in enumerate(columnas, 1):
        celda = detalle.cell(row=1, column=col, value=texto)
        celda.fill, celda.font = cabecera, blanco
    for i, t in enumerate(d["tickets"], 2):
        valores = [t["codigo"], t["empresa"], t["asunto"], t["categoria"], t["prioridad"], t["estado"],
                   t["solicitante"], t["area"] or "-", t["tecnico"] or "Sin asignar",
                   t["created_at"].strftime("%d/%m/%Y %H:%M"), t["fecha_limite"].strftime("%d/%m/%Y %H:%M"),
                   t["sla_texto"]]
        for col, valor in enumerate(valores, 1):
            detalle.cell(row=i, column=col, value=valor)
    anchos = [12, 18, 42, 18, 10, 12, 22, 18, 20, 17, 17, 22]
    for col, ancho in enumerate(anchos, 1):
        detalle.column_dimensions[get_column_letter(col)].width = ancho
    detalle.freeze_panes = "A2"
    for celda in detalle[1]:
        celda.alignment = Alignment(vertical="center")

    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"reporte_tickets_{ahora():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------------ PDF
@dashboard_bp.get("/dashboard/exportar/pdf")
@login_requerido
def exportar_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    filtros = leer_filtros()
    d = datos_dashboard(filtros)
    primario = colors.HexColor(g.usuario.get("color_primario") or "#3D4A8A")
    estilos = getSampleStyleSheet()
    chico = estilos["BodyText"].clone("chico", fontSize=8, leading=10)

    def tabla(filas, anchos):
        t = Table(filas, colWidths=anchos, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), primario),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F1EC")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D9D4C7")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return t

    salida = io.BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title="Reporte de tickets")
    titulo = g.usuario.get("titulo_sistema") or "Mesa de Ayuda TI"
    partes = [
        Paragraph(f"Reporte de tickets — {titulo}", estilos["Title"]),
        Paragraph(texto_filtros(filtros), estilos["Normal"]),
        Paragraph(f"Generado el {ahora():%d/%m/%Y %H:%M} por {g.usuario['nombre']}", estilos["Normal"]),
        Spacer(1, 0.4 * cm),
    ]
    resumen = [["Total", *ESTADOS, "Verde", "Amarillo", "Rojo", "% en plazo"],
               [d["total"], *[d["por_estado"][e] for e in ESTADOS], d["sla"]["verde"], d["sla"]["amarillo"],
                d["sla"]["rojo"], f"{d['cumplimiento']}%"]]
    partes += [tabla(resumen, [2.4 * cm] * 9), Spacer(1, 0.4 * cm)]
    resp = [["Técnico", "Asignados", "En proceso", "Resueltos", "A tiempo"]] + [
        [r["tecnico"], r["asignados"], r["en_proceso"], r["resueltos"],
         f"{r['porcentaje']}%" if r["porcentaje"] is not None else "-"] for r in d["responsables"]]
    partes += [Paragraph("Por responsable", estilos["Heading3"]), tabla(resp, [7 * cm, 3 * cm, 3 * cm, 3 * cm, 3 * cm]),
               Spacer(1, 0.4 * cm), Paragraph("Detalle de tickets", estilos["Heading3"])]
    detalle = [["Código", "Asunto", "Empresa", "Prioridad", "Estado", "Solicitante", "Técnico", "Creado", "SLA"]] + [
        [t["codigo"], Paragraph(t["asunto"], chico), t["empresa"], t["prioridad"], t["estado"],
         Paragraph(t["solicitante"], chico), Paragraph(t["tecnico"] or "Sin asignar", chico),
         t["created_at"].strftime("%d/%m/%Y %H:%M"), t["sla_texto"]] for t in d["tickets"]]
    partes.append(tabla(detalle, [2.1 * cm, 6.2 * cm, 3.2 * cm, 1.8 * cm, 2 * cm, 3.4 * cm, 3.4 * cm, 2.8 * cm, 3 * cm]))
    doc.build(partes)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"reporte_tickets_{ahora():%Y%m%d_%H%M}.pdf",
                     mimetype="application/pdf")


# ------------------------------------------------------------------ API
@dashboard_bp.get("/api/dashboard")
@login_requerido
def api_dashboard():
    filtros = leer_filtros()
    d = datos_dashboard(filtros)
    d.pop("tickets")
    return jsonify(status=1, filtros={k: (v.strftime("%Y-%m-%d") if isinstance(v, datetime) else v)
                                      for k, v in filtros.items()}, data=d)
