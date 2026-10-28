"""
MÓDULO 6 — Exportación de reportes con los mismos filtros de la pantalla Reportes.

  GET /reportes/exportar/excel   -> .xlsx (openpyxl)
  GET /reportes/exportar/pdf     -> .pdf  (ReportLab)
  GET /api/reportes              -> resumen JSON (para Postman)

Filtros (query string): desde, hasta (AAAA-MM-DD), area, categoria, prioridad,
estado (Abierto, En proceso, Resuelto, Cerrado o Vencido), tecnico, q.
"""
import io
from datetime import datetime, timedelta

from flask import Blueprint, g, jsonify, request, send_file, session

from auth import ROL_ADMIN, ROL_SUPER, ROL_TECNICO, empresa_actual, roles_requeridos
from db import consultar
from utils import ESTADOS, ahora
from vistas import V_TICKETS, aplicar_sla_lista

reportes_bp = Blueprint("reportes_export", __name__)


def _fecha(texto):
    try:
        return datetime.strptime(texto, "%Y-%m-%d")
    except (TypeError, ValueError):
        return None


def filtros_de_la_peticion():
    a = request.args
    return {k: (a.get(k) or "").strip() for k in ("desde", "hasta", "area", "categoria", "prioridad", "estado", "tecnico", "q")}


def tickets_filtrados(empresa_id, f):
    condiciones, params = ["v.empresa_id = %s"], [empresa_id]
    if _fecha(f["desde"]):
        condiciones.append("v.created_at >= %s"); params.append(_fecha(f["desde"]))
    if _fecha(f["hasta"]):
        condiciones.append("v.created_at < %s"); params.append(_fecha(f["hasta"]) + timedelta(days=1))
    for campo in ("area", "categoria", "prioridad", "tecnico"):
        if f[campo]:
            condiciones.append(f"v.{campo} = %s"); params.append(f[campo])
    if f["estado"] in ESTADOS:
        condiciones.append("v.estado = %s"); params.append(f["estado"])
    if f["q"]:
        condiciones.append("(v.codigo LIKE %s OR v.asunto LIKE %s OR v.solicitante LIKE %s OR v.tecnico LIKE %s)")
        params += [f"%{f['q']}%"] * 4
    filas = consultar(f"SELECT * FROM {V_TICKETS} AS v WHERE {' AND '.join(condiciones)} ORDER BY v.created_at DESC", params)
    filas = aplicar_sla_lista(filas)
    if f["estado"] == "Vencido":
        filas = [t for t in filas if t["semaforo"] == "rojo"]
    return filas


def resumen(tickets):
    por_estado = {e: sum(1 for t in tickets if t["estado"] == e) for e in ESTADOS}
    vencidos = sum(1 for t in tickets if t["semaforo"] == "rojo")
    terminados = [t for t in tickets if t["estado"] in ("Resuelto", "Cerrado")]
    a_tiempo = sum(1 for t in terminados if t["semaforo"] == "cumplido")
    def agrupar(campo):
        datos = {}
        for t in tickets:
            datos[t[campo] or "—"] = datos.get(t[campo] or "—", 0) + 1
        return sorted(datos.items(), key=lambda x: -x[1])
    return {
        "total": len(tickets), "por_estado": por_estado, "vencidos": vencidos,
        "cumplimiento": round(100 * a_tiempo / len(terminados)) if terminados else None,
        "por_prioridad": agrupar("prioridad"), "por_categoria": agrupar("categoria"),
        "por_tecnico": agrupar("tecnico"), "por_area": agrupar("area"),
    }


def texto_filtros(f):
    partes = []
    if f["desde"] or f["hasta"]:
        partes.append(f"Fechas: {f['desde'] or 'inicio'} a {f['hasta'] or 'hoy'}")
    for campo, nombre in (("area", "Área"), ("categoria", "Categoría"), ("prioridad", "Prioridad"),
                          ("estado", "Estado"), ("tecnico", "Técnico"), ("q", "Búsqueda")):
        if f[campo]:
            partes.append(f"{nombre}: {f[campo]}")
    return " | ".join(partes) or "Sin filtros (todos los tickets)"


FILAS = ("Código", "Asunto", "Solicitante", "Área", "Categoría", "Prioridad", "Estado", "Técnico",
         "Creado", "Fecha límite", "Resuelto", "SLA")


def fila(t):
    f = lambda d: d.strftime("%d/%m/%Y %H:%M") if d else "—"  # noqa: E731
    return [t["codigo"], t["asunto"], t["solicitante"], t["area"] or "—", t["categoria"], t["prioridad"],
            t["estado"], t["tecnico"], f(t["created_at"]), f(t["fecha_limite"]), f(t["fecha_resolucion"]), t["sla_texto"]]


@reportes_bp.get("/reportes/exportar/excel")
@roles_requeridos(ROL_SUPER, ROL_ADMIN, ROL_TECNICO)
def exportar_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    f = filtros_de_la_peticion()
    tickets = tickets_filtrados(empresa_actual(), f)
    r = resumen(tickets)
    color = (session.get("empresa_color") or "#0077B6").lstrip("#")
    cab, blanco = PatternFill("solid", fgColor=color), Font(bold=True, color="FFFFFF")

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    hoja["A1"] = f"Reporte de tickets - {session.get('empresa_nombre', '')}"
    hoja["A1"].font = Font(bold=True, size=14)
    hoja["A2"] = texto_filtros(f)
    hoja["A3"] = f"Generado el {ahora():%d/%m/%Y %H:%M} por {session.get('nombre', '')}"
    fila_actual = 5

    def bloque(titulo, encabezados, filas):
        nonlocal fila_actual
        hoja.cell(row=fila_actual, column=1, value=titulo).font = Font(bold=True, size=12)
        fila_actual += 1
        for col, texto in enumerate(encabezados, 1):
            celda = hoja.cell(row=fila_actual, column=col, value=texto)
            celda.fill, celda.font = cab, blanco
        for valores in filas:
            fila_actual += 1
            for col, valor in enumerate(valores, 1):
                hoja.cell(row=fila_actual, column=col, value=valor)
        fila_actual += 2

    bloque("Indicadores", ["Indicador", "Valor"],
           [("Total de tickets", r["total"]), *[(e, n) for e, n in r["por_estado"].items()],
            ("Vencidos (SLA)", r["vencidos"]),
            ("Cumplimiento SLA de resueltos", f"{r['cumplimiento']}%" if r["cumplimiento"] is not None else "—")])
    bloque("Por prioridad", ["Prioridad", "Tickets"], r["por_prioridad"])
    bloque("Por categoría", ["Categoría", "Tickets"], r["por_categoria"])
    bloque("Por técnico", ["Técnico", "Tickets"], r["por_tecnico"])
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 14

    detalle = libro.create_sheet("Tickets")
    for col, texto in enumerate(FILAS, 1):
        celda = detalle.cell(row=1, column=col, value=texto)
        celda.fill, celda.font = cab, blanco
    for i, t in enumerate(tickets, 2):
        for col, valor in enumerate(fila(t), 1):
            detalle.cell(row=i, column=col, value=valor)
    for col, ancho in enumerate((12, 40, 22, 18, 18, 10, 12, 20, 17, 17, 17, 22), 1):
        detalle.column_dimensions[get_column_letter(col)].width = ancho
    detalle.freeze_panes = "A2"
    detalle.auto_filter.ref = f"A1:L{max(len(tickets) + 1, 1)}"

    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"reporte_tickets_{ahora():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@reportes_bp.get("/reportes/exportar/pdf")
@roles_requeridos(ROL_SUPER, ROL_ADMIN, ROL_TECNICO)
def exportar_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    f = filtros_de_la_peticion()
    tickets = tickets_filtrados(empresa_actual(), f)
    r = resumen(tickets)
    primario = colors.HexColor(session.get("empresa_color") or "#0077B6")
    estilos = getSampleStyleSheet()
    chico = estilos["BodyText"].clone("chico", fontSize=7.5, leading=9.5)

    def tabla(filas, anchos):
        t = Table(filas, colWidths=anchos, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), primario), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF4FB")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D8EC")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return t

    salida = io.BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title="Reporte de tickets")
    partes = [
        Paragraph(f"Reporte de tickets — {session.get('empresa_titulo') or session.get('empresa_nombre', '')}", estilos["Title"]),
        Paragraph(texto_filtros(f), estilos["Normal"]),
        Paragraph(f"Generado el {ahora():%d/%m/%Y %H:%M} por {session.get('nombre', '')}", estilos["Normal"]),
        Spacer(1, 0.4 * cm),
        tabla([["Total", *ESTADOS, "Vencidos", "Cumplimiento SLA"],
               [r["total"], *r["por_estado"].values(), r["vencidos"],
                f"{r['cumplimiento']}%" if r["cumplimiento"] is not None else "—"]], [3 * cm] * 7),
        Spacer(1, 0.4 * cm),
        tabla([[Paragraph(f"<b>{c}</b>", chico) for c in FILAS]] +
              [[Paragraph(str(v), chico) for v in fila(t)] for t in tickets],
              [1.8 * cm, 5 * cm, 2.8 * cm, 2.2 * cm, 2.2 * cm, 1.5 * cm, 1.7 * cm, 2.6 * cm, 2.2 * cm, 2.2 * cm, 2.2 * cm, 2.4 * cm]),
    ]
    doc.build(partes)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"reporte_tickets_{ahora():%Y%m%d_%H%M}.pdf",
                     mimetype="application/pdf")


@reportes_bp.get("/api/reportes")
@roles_requeridos(ROL_SUPER, ROL_ADMIN, ROL_TECNICO)
def api_reportes():
    f = filtros_de_la_peticion()
    return jsonify(status=1, filtros=f, data=resumen(tickets_filtrados(empresa_actual(), f)))


@reportes_bp.route("/api/sla/configuracion", methods=["GET", "PUT", "POST"])
@roles_requeridos(ROL_SUPER, ROL_ADMIN)
def configuracion_sla():
    """Plazos por prioridad (tabla prioridades). Son globales: solo el superusuario los cambia."""
    from db import conexion
    if request.method == "GET":
        return jsonify(prioridades=consultar("SELECT id, nombre, sla_horas FROM prioridades ORDER BY nivel DESC"))
    if g.usuario_token["rol"] != ROL_SUPER:
        return jsonify(error="Los plazos son iguales para todas las empresas: solo el superusuario puede cambiarlos."), 403
    datos = request.get_json(silent=True) or {}
    cambios = datos.get("prioridades") or {}
    con = conexion()
    with con.cursor() as cur:
        for nombre, horas in cambios.items():
            try:
                horas = int(horas)
            except (TypeError, ValueError):
                return jsonify(error=f"Horas no válidas para {nombre}."), 400
            if not 1 <= horas <= 720:
                return jsonify(error="Las horas deben estar entre 1 y 720."), 400
            cur.execute("UPDATE prioridades SET sla_horas = %s WHERE nombre = %s", (horas, nombre))
    con.commit()
    return jsonify(mensaje="Plazos guardados. Se aplican a los tickets nuevos o a los que cambien de prioridad.",
                   prioridades=consultar("SELECT id, nombre, sla_horas FROM prioridades ORDER BY nivel DESC"))
