"""Reportes exportables: Excel con openpyxl y PDF con ReportLab (módulo 6)."""
from datetime import timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .servicios import zona_local

COLUMNAS = ["Código", "Asunto", "Categoría", "Prioridad", "Estado", "Solicitante", "Empresa",
            "Técnico", "Creado", "Tiempo (SLA)"]
HORTENSIA = "D9DEF5"
HORTENSIA_PROFUNDO = "3F4C8C"


def _fecha_local(dt):
    return dt.replace(tzinfo=timezone.utc).astimezone(zona_local()).strftime("%d/%m/%Y %H:%M")


def _filas(tickets):
    return [[t.codigo, t.asunto, t.categoria, t.prioridad, t.estado, t.solicitante.nombre, t.empresa,
             t.tecnico.nombre if t.tecnico else "Sin asignar", _fecha_local(t.created_at), t.sla["texto"]]
            for t in tickets]


def descripcion_filtros(filtros, empresa):
    partes = [f"Del {filtros['desde']} al {filtros['hasta']}"]
    partes.append(empresa or "Todas las empresas")
    if filtros.get("area"):
        partes.append(f"Área: {filtros['area']}")
    return " · ".join(partes)


def generar_excel(tickets, subtitulo):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tickets"

    hoja["A1"] = "Reporte de tickets · Mesa de Ayuda TI"
    hoja["A1"].font = Font(bold=True, size=14, color=HORTENSIA_PROFUNDO)
    hoja["A2"] = subtitulo
    hoja["A2"].font = Font(italic=True, color="6B707A")

    hoja.append([])
    hoja.append(COLUMNAS)
    for celda in hoja[4]:
        celda.font = Font(bold=True, color=HORTENSIA_PROFUNDO)
        celda.fill = PatternFill("solid", fgColor=HORTENSIA)
        celda.alignment = Alignment(vertical="center")

    for fila in _filas(tickets):
        hoja.append(fila)

    anchos = [10, 40, 18, 10, 12, 22, 20, 20, 17, 22]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A5"
    if tickets:
        hoja.auto_filter.ref = f"A4:{get_column_letter(len(COLUMNAS))}{hoja.max_row}"

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def generar_pdf(tickets, subtitulo, datos_resumen):
    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
                                  topMargin=14 * mm, bottomMargin=14 * mm, title="Reporte de tickets")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], textColor=colors.HexColor("#3F4C8C"), alignment=0)
    sub = ParagraphStyle("sub", parent=estilos["Normal"], textColor=colors.HexColor("#6B707A"))
    celda = ParagraphStyle("celda", parent=estilos["Normal"], fontSize=8, leading=10)

    elementos = [Paragraph("Reporte de tickets · Mesa de Ayuda TI", titulo), Paragraph(subtitulo, sub), Spacer(1, 6 * mm)]

    estados = datos_resumen["por_estado"]
    resumen_tabla = Table([["Total", "Abiertos", "En proceso", "Resueltos", "Cerrados", "Dentro del plazo"],
                           [datos_resumen["total"], estados["Abierto"], estados["En proceso"], estados["Resuelto"],
                            estados["Cerrado"],
                            f"{datos_resumen['sla_porcentaje']}%" if datos_resumen["sla_porcentaje"] is not None else "—"]])
    resumen_tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9DEF5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#3F4C8C")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E8E3DA")),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elementos += [resumen_tabla, Spacer(1, 6 * mm)]

    if tickets:
        filas = [COLUMNAS] + [[Paragraph(str(v), celda) for v in fila] for fila in _filas(tickets)]
        tabla = Table(filas, repeatRows=1,
                      colWidths=[18 * mm, 58 * mm, 26 * mm, 16 * mm, 20 * mm, 28 * mm, 28 * mm, 28 * mm, 25 * mm, 28 * mm])
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3F4C8C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F4EE")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E8E3DA")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("No hay tickets en el rango elegido.", sub))

    documento.build(elementos)
    salida.seek(0)
    return salida
